"""Documentos de hipótesis controlados, según la norma AP-DOC-002 v01
("Hypothesis Document Control", borrador del 17 de septiembre de 2026, de
Monica Duarte, pendiente de revisión de la dirección).

Qué pide la norma a cada documento de hipótesis:

- Un código propio, correlativo y que no se reutiliza: AP-HYP-001, AP-HYP-002...
  (AP de Alzheimer Project, HYP de hipótesis, tres cifras).
- Tipo DOC, aunque el contenido sea científico.
- Versión con dos cifras enteras: v01, v02...
- Cabecera: [nombre oficial corto] | DOC | AP-HYP-NNN | vNN
- Pie: Confidential | Alzheimer Project | AI Robotix | MMM-DD-YYYY | [iniciales
  del responsable] | Page X of Y, con la numeración como campo automático, nunca
  escrita a mano.
- Una lista de ocho comprobaciones antes de darlo por controlado.

En ROSA2018 el documento de hipótesis es el dossier para el laboratorio
(rosa/dossier.py), que ya se guarda como artefacto con versiones. Emitirlo como
documento controlado es un acto aparte y deliberado (`emitir`): el código lo
asigna siempre el servidor, la primera vez, y queda en la hipótesis para
siempre; cada emisión nueva sube la versión y guarda con ella lo que llevaba
(nombre, fecha, iniciales, cabecera, pie y el resultado de las ocho
comprobaciones), de modo que la v01 se descarga igual aunque luego cambie el
nombre corto o las iniciales. El Word se construye al descargarlo a partir de
esa versión y del contenido del dossier que se emitió (`docx`).

La norma es un borrador: todo lo que fija está aquí, en un solo sitio.
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any

SOP = "AP-DOC-002 v01"
PREFIJO = "AP"
AREA = "HYP"
TIPO = "DOC"
CONFIDENCIALIDAD = "Confidential"
PROYECTO = "Alzheimer Project"
ORGANIZACION = "AI Robotix"

# Las iniciales del responsable que van en el pie. La norma pide las del dueño
# del documento (su propio ejemplo es MD, una persona). Emir pidió "AP" el 23 de
# septiembre de 2026, a confirmar con Monica: si cambia, las versiones ya
# emitidas conservan las suyas y las nuevas llevan las nuevas.
INICIALES_RESPONSABLE = "AP"

# La fecha del pie, en el formato de la norma (Sep-17-2026). Los meses van en
# inglés como en su ejemplo, y fijos: no dependen del idioma de la máquina.
_MESES = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")

RE_ID = re.compile(rf"^{PREFIJO}-{AREA}-\d{{3,}}$")
RE_VERSION = re.compile(r"^v\d{2,}$")
RE_FECHA = re.compile(r"^(" + "|".join(_MESES) + r")-\d{2}-\d{4}$")
RE_INICIALES = re.compile(r"^[A-Z]{2,4}$")
RE_CABECERA = re.compile(rf"^[^|]+ \| {TIPO} \| {PREFIJO}-{AREA}-\d{{3,}} \| v\d{{2,}}$")
RE_PIE = re.compile(rf"^{CONFIDENCIALIDAD} \| {PROYECTO} \| {ORGANIZACION} \| (" + "|".join(_MESES) + r")-\d{2}-\d{4} \| [A-Z]{2,4} \| Page X of Y$")

MAX_NOMBRE = 80


def id_documento(n: int) -> str:
    return f"{PREFIJO}-{AREA}-{n:03d}"


def version_documento(n: int) -> str:
    return f"v{n:02d}"


def fecha_documento(ms: int) -> str:
    d = datetime.fromtimestamp(ms / 1000)
    return f"{_MESES[d.month - 1]}-{d.day:02d}-{d.year}"


def cabecera(nombre: str, id_doc: str, version: str) -> str:
    return f"{nombre} | {TIPO} | {id_doc} | {version}"


def pie(fecha: str, iniciales: str) -> str:
    """El pie tal como lo escribe la norma; en el Word, "X" e "Y" son los
    campos automáticos de página y de número de páginas."""
    return f"{CONFIDENCIALIDAD} | {PROYECTO} | {ORGANIZACION} | {fecha} | {iniciales} | Page X of Y"


def limpiar_nombre(nombre: Any) -> str:
    """El nombre corto tal como puede ir en la cabecera: sin la barra vertical,
    que partiría la cabecera, sin comillas ni punto final, en una línea."""
    t = re.sub(r"\s+", " ", str(nombre or "").replace("|", " ")).strip().strip("\"'«»“”").strip().rstrip(".").strip()
    return t[:MAX_NOMBRE].strip()


def numero_de(id_doc: Any) -> int:
    m = re.match(rf"^{PREFIJO}-{AREA}-(\d+)$", str(id_doc or ""))
    return int(m.group(1)) if m else 0


def siguiente_numero(e: dict[str, Any]) -> int:
    """El siguiente número libre. Se guarda el último asignado y además se mira
    el mayor que haya en las hipótesis: aunque el contador se perdiera, un
    código ya usado no vuelve a salir."""
    ultimo = int((e.get("documentosControlados") or {}).get("ultimo") or 0)
    usados = [numero_de((h.get("documentoControlado") or {}).get("id")) for h in e.get("hipotesis", []) if isinstance(h, dict)]
    return max([ultimo, *usados]) + 1


# Las ocho comprobaciones de la sección 8 de la norma, con su texto original.
COMPROBACIONES = (
    (1, "Header follows the approved format", "La cabecera sigue el formato aprobado"),
    (2, "Footer follows the approved format", "El pie sigue el formato aprobado"),
    (3, "Document type is DOC", "El tipo de documento es DOC"),
    (4, "Area code is HYP", "El código de área es HYP"),
    (5, "Document ID follows AP-HYP-[Sequential Number]", "El código sigue AP-HYP-[número correlativo]"),
    (6, "Version is assigned", "La versión está asignada"),
    (7, "Date is included in the footer", "La fecha está en el pie"),
    (8, "Owner initials are included in the footer", "Las iniciales del responsable están en el pie"),
)


def comprobar(id_doc: str, v: dict[str, Any]) -> list[dict[str, Any]]:
    """La lista de la sección 8, por regla, sobre lo que se va a imprimir. Como
    lo genera ROSA2018, debería salir siempre entera; está para que un cambio de
    configuración (unas iniciales vacías, un nombre con barra) no emita un
    documento que no cumple sin que nadie lo vea."""
    cab, pi = str(v.get("cabecera") or ""), str(v.get("pie") or "")
    ok = {
        1: bool(RE_CABECERA.match(cab)) and bool(limpiar_nombre(v.get("nombre"))),
        2: bool(RE_PIE.match(pi)),
        3: f" | {TIPO} | " in cab,
        4: f"-{AREA}-" in str(id_doc),
        5: bool(RE_ID.match(str(id_doc))),
        6: bool(RE_VERSION.match(str(v.get("version") or ""))),
        7: bool(RE_FECHA.match(str(v.get("fecha") or ""))) and f" | {v.get('fecha')} | " in pi,
        8: bool(RE_INICIALES.match(str(v.get("iniciales") or ""))) and f" | {v.get('iniciales')} | Page" in pi,
    }
    return [{"n": n, "norma": en, "texto": es, "ok": ok[n]} for n, en, es in COMPROBACIONES]


def emitir(e: dict[str, Any], h: dict[str, Any], nombre_corto: Any, quien: str, ahora: int) -> dict[str, Any] | str:
    """Emite como documento controlado la última versión del dossier de la
    hipótesis. Devuelve la versión emitida o, si no se puede, el motivo."""
    art_id = h.get("dossierArtefactoId")
    art = next((a for a in e.get("artefactos", []) if isinstance(a, dict) and a.get("id") == art_id), None) if art_id else None
    if not art or not art.get("versiones"):
        return "La hipótesis no tiene dossier: genéralo primero."
    n_art = len(art["versiones"])
    doc = h.get("documentoControlado") if isinstance(h.get("documentoControlado"), dict) else None
    versiones = doc.get("versiones") if doc and isinstance(doc.get("versiones"), list) else []
    if versiones and versiones[-1].get("artefactoId") == art_id and versiones[-1].get("versionArtefacto") == n_art:
        return f"Esa versión del dossier ya está emitida como {versiones[-1].get('version')}: vuelve a generar el dossier si ha cambiado."
    nombre = limpiar_nombre(nombre_corto) or limpiar_nombre((doc or {}).get("nombreCorto")) or limpiar_nombre(h.get("nombreCorto"))
    if not nombre:
        return "Falta el nombre corto de la cabecera."
    if doc is None:
        n = siguiente_numero(e)
        e.setdefault("documentosControlados", {})["ultimo"] = n
        doc = {"id": id_documento(n), "nombreCorto": nombre, "versiones": []}
        h["documentoControlado"] = doc
        versiones = doc["versiones"]
    else:
        doc["nombreCorto"] = nombre
        doc["versiones"] = versiones
    fecha = fecha_documento(ahora)
    version = version_documento(len(versiones) + 1)
    v: dict[str, Any] = {
        "version": version,
        "emitidaEn": ahora,
        "fecha": fecha,
        "iniciales": INICIALES_RESPONSABLE,
        "nombre": nombre,
        "cabecera": cabecera(nombre, doc["id"], version),
        "pie": pie(fecha, INICIALES_RESPONSABLE),
        "artefactoId": art_id,
        "versionArtefacto": n_art,
        "quien": quien,
        "norma": SOP,
    }
    v["comprobaciones"] = comprobar(doc["id"], v)
    v["controlado"] = all(c["ok"] for c in v["comprobaciones"])
    versiones.append(v)
    return v


def para_descargar(e: dict[str, Any], art_id: str, n: int, ahora: int) -> tuple[str, str | None, dict[str, Any]] | None:
    """Lo que hace falta para descargar la versión `n` de un dossier: su
    contenido, el código y la versión emitida si esa versión del dossier se
    emitió, o un borrador si no. None si no existe."""
    art = next((a for a in e.get("artefactos", []) if isinstance(a, dict) and a.get("id") == art_id), None)
    if not art or art.get("tipo") != "dossier" or not 1 <= n <= len(art.get("versiones") or []):
        return None
    contenido = str(art["versiones"][n - 1].get("contenido") or "")
    h = next((x for x in e.get("hipotesis", []) if isinstance(x, dict) and x.get("dossierArtefactoId") == art_id), None) or {}
    doc = h.get("documentoControlado") if isinstance(h.get("documentoControlado"), dict) else None
    v = next((x for x in reversed((doc or {}).get("versiones") or []) if isinstance(x, dict) and x.get("artefactoId") == art_id and x.get("versionArtefacto") == n), None)
    if v and doc:
        return contenido, str(doc["id"]), dict(v)
    return contenido, None, borrador(str((doc or {}).get("nombreCorto") or h.get("nombreCorto") or h.get("titulo") or ""), ahora)


def version_emitida(h: dict[str, Any], version: str) -> dict[str, Any] | None:
    doc = h.get("documentoControlado") if isinstance(h.get("documentoControlado"), dict) else None
    return next((v for v in (doc or {}).get("versiones") or [] if isinstance(v, dict) and v.get("version") == version), None)


# -- El Word ----------------------------------------------------------------------
#
# El diseño es el de la propia norma: la plantilla rosa/plantillas/
# documento_hipotesis.docx se sacó del Word de Monica Duarte (25 de septiembre
# de 2026) con su cabecera (el árbol del Alzheimer Project y la línea de control
# a la derecha), su pie (con PAGE y NUMPAGES como campos de Word) y un párrafo de
# muestra de cada tipo: título, subtítulo, línea de datos, sección, texto,
# recuadro y viñeta. Cada párrafo del documento es una copia de su muestra con
# otro texto, así que la letra, los colores, los tamaños, el sombreado lila y los
# espaciados son los suyos, no una imitación.

PLANTILLA = Path(__file__).resolve().parent / "plantillas" / "documento_hipotesis.docx"
MAX_NOMBRE_BORRADOR = 50
MUESTRAS = ("titulo", "subtitulo", "vacio", "dato", "seccion", "texto", "recuadro", "vineta")


def borrador(nombre: str, ahora: int) -> dict[str, Any]:
    """La cabecera y el pie de un dossier que todavía no se ha emitido como
    documento controlado: el mismo diseño, pero sin código AP-HYP, que solo lo
    da la emisión, y dicho en la cabecera para que nadie lo tome por uno."""
    fecha = fecha_documento(ahora)
    nombre = limpiar_nombre(nombre) or "Dossier para el laboratorio"
    # Si todavía no hay nombre corto llega el título entero: se corta por una
    # palabra entera, para que la cabecera quepa en una línea.
    if len(nombre) > MAX_NOMBRE_BORRADOR:
        nombre = nombre[:MAX_NOMBRE_BORRADOR].rsplit(" ", 1)[0].rstrip(" ,;:-–") + "…"
    return {"version": "borrador", "fecha": fecha, "iniciales": INICIALES_RESPONSABLE, "nombre": nombre, "cabecera": f"{nombre} | {TIPO} | sin código: no emitido | borrador", "pie": pie(fecha, INICIALES_RESPONSABLE), "emitido": False}


def _con_texto(muestra: Any, texto: str) -> Any:
    """Una copia del párrafo de muestra con otro texto. Las **negritas** del
    Markdown del dossier se conservan como trozos en negrita con el mismo
    formato de la muestra."""
    import copy

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    nuevo = copy.deepcopy(muestra)
    runs = nuevo.findall(qn("w:r"))
    base = copy.deepcopy(runs[0]) if runs else None
    for r in runs:
        nuevo.remove(r)
    if base is None:
        return nuevo
    for t in base.findall(qn("w:t")):
        base.remove(t)
    for i, trozo in enumerate(re.split(r"\*\*", texto)):
        if not trozo:
            continue
        r = copy.deepcopy(base)
        if i % 2 == 1:
            rpr = r.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                r.insert(0, rpr)
            if rpr.find(qn("w:b")) is None:
                rpr.append(OxmlElement("w:b"))
        # Arial explícito, como el PDF de la norma: el Word lo trae como fuente por
        # defecto del documento, y hay visores (la vista previa de macOS) que no la
        # aplican y ponen Times.
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            r.insert(0, rpr)
        if rpr.find(qn("w:rFonts")) is None:
            fuentes = OxmlElement("w:rFonts")
            for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fuentes.set(qn(k), "Arial")
            rpr.insert(0, fuentes)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = trozo
        r.append(t)
        nuevo.append(r)
    return nuevo


def docx(contenido: str, id_doc: str | None, v: dict[str, Any]) -> bytes:
    """El dossier en Word con el diseño de la norma AP-DOC-002. `v` es una
    versión emitida (con su código `id_doc`) o un `borrador`."""
    from docx import Document

    d = Document(str(PLANTILLA))
    sec = d.sections[0]
    # Cabecera y pie: se cambia el texto y se conservan logo y campos.
    cab = next(p for p in sec.header.paragraphs if "{{cabecera}}" in p.text)
    cab.runs[0].text = v["cabecera"]
    pie_p = sec.footer.paragraphs[0]
    pie_p.runs[0].text = v["pie"].split("Page X of Y")[0] + "Page "
    for r in (*cab.runs, *pie_p.runs):
        r.font.name = "Arial"

    cuerpo = d.element.body
    muestras = {}
    for p in list(d.paragraphs):
        clave = p.text.strip().strip("{}")
        if clave in MUESTRAS:
            muestras[clave] = p._p
            cuerpo.remove(p._p)
    fin = cuerpo[-1] if len(cuerpo) and cuerpo[-1].tag.endswith("sectPr") else None

    def poner(tipo: str, texto: str = "") -> None:
        el = _con_texto(muestras[tipo], texto) if texto else _con_texto(muestras["vacio"], "")
        if fin is not None:
            fin.addprevious(el)
        else:
            cuerpo.append(el)

    lineas = [x.rstrip() for x in para_lector(contenido).splitlines()]
    titulo = next((x[2:].strip() for x in lineas if x.startswith("# ")), "Dossier para el laboratorio")
    emitido = v.get("emitido", True)
    poner("titulo", titulo)
    poner("subtitulo", "Dossier para el laboratorio, generado por ROSA2018")
    poner("vacio")
    datos = [
        f"Document Status: Draft{'' if emitido else ' (no emitido como documento controlado)'}",
        f"Document ID: {id_doc if emitido and id_doc else 'sin asignar (se da al emitir)'}",
        f"Version: {v['version'][1:] if emitido else 'borrador'}",
        f"Date: {v['fecha']}",
        f"Owner Initials: {v['iniciales']}",
    ]
    for x in datos:
        poner("dato", x)
    poner("vacio")
    vio_titulo = False
    for x in lineas:
        if not x.strip():
            continue
        if x.startswith("# ") and not vio_titulo:
            vio_titulo = True  # ya es el título del documento
            continue
        if x.startswith("#"):
            poner("vacio")
            poner("seccion", x.lstrip("#").strip())
        elif re.match(r"^\s*[-*] ", x):
            poner("vineta", "- " + re.sub(r"^\s*[-*] ", "", x))
        else:
            poner("texto", x)

    salida = io.BytesIO()
    d.save(salida)
    return salida.getvalue()


# -- Lo que el lector no necesita ------------------------------------------------------
#
# El dossier guardado lleva todo lo que hace falta para auditarlo: ids internos,
# el commit, el Elo, las comprobaciones una a una del auditor, los resultados en
# bruto del análisis, la tabla de conectores. Emir (25 de septiembre de 2026):
# "esos datos crudos no deberían verse en el pdf". El documento que se descarga
# es para quien lo lee en el laboratorio; el artefacto guardado no se toca.

_SECCIONES_FUERA = ("qué dicen las bases de la diana", "que dicen las bases de la diana", "revisión del registro", "revision del registro")
_LINEAS_FUERA = re.compile(
    r"^\s*(ROSA2018: commit|Nivel de autonom[ií]a|Novedad:|Resultados:|Baseline:|Auditor[ií]a\b|Perfil de evidencia por diana|Capa \| Estado)"
    r"|^\s*-\s*[a-z]+(?:_[a-z0-9]+)*:\s*(pasa|falla|no_aplica|no_comprobable|no aplica)\b"
    r"|:\s*(no declarad[oa]s?|sin criterio)\s*\.?$"
    r"|\|.*\|.*\|"
    r"|^\s*LM Response|^\s*[{}\[\]]|^\s*\"[a-z_]+\"\s*:"
    r"|El juez no respondi[oó]",
    re.IGNORECASE,
)
_BLOQUEO_LEGIBLE = {
    "revision_registro_abierta": "La revisión del registro tiene hallazgos graves sin atender",
    "dependencia_pendiente": "Algo de lo que depende cambió y está pendiente de revisar",
    "trazabilidad_insuficiente": "Trazabilidad insuficiente: no hay afirmaciones sostenidas o alguna está bloqueada",
    "sin_experimento_interpretable": "Sin experimento interpretable: faltan los criterios de confirmación o refutación",
    "descartada_por_killer": "Descartada en este contexto",
    "fuente_retractada": "Depende de una fuente retractada",
    "datos_no_autorizados": "Datos no autorizados para su uso con IA",
    "analisis_invalido": "Análisis inválido según el auditor",
}
_ESTADO_SUPUESTO = {"sin_evidencia": "sin evidencia", "respaldado": "respaldado", "plausible": "plausible", "contradicho": "contradicho"}


def _limpiar_linea(x: str) -> str:
    t = x
    # Ids internos y huellas.
    t = re.sub(r"Hip[oó]tesis\s+hip-[\w-]+,\s*versi[oó]n\s+(\d+)\.", r"Versión \1 de la hipótesis.", t)
    t = re.sub(r"\s*\((?:artefacto|plan|datos)\s+[\w-]+\)", "", t)
    t = re.sub(r"\[An[aá]lisis in silico[^\]]*\]", "[Análisis in silico]", t)
    t = re.sub(r"\[(Datos (?:del laboratorio|de prueba, SINT[EÉ]TICOS)):\s*[^,\]]+,\s*", r"[\1, ", t)
    t = re.sub(r"\b(?:hip|art|run|cor|inv|af|it|plan|ev)-[a-z0-9]+(?:-[a-z0-9]+)*\b", "", t)
    t = re.sub(r"\b(?:datos\s+)?sha256\s+[0-9a-f]+\b|\b[0-9a-f]{12,64}\b", "", t)
    # Lo interno del torneo y del Killer.
    t = re.sub(r"\s*Elo\s+\d+\s+tras\s+\d+\s+partidos\.", "", t)
    t = re.sub(r"Decisi[oó]n del Killer sobre esta versi[oó]n", "Revisión crítica de esta versión", t)
    t = re.sub(r"\s*·\s*killer_\d+", "", t)
    t = re.sub(r"\s*·\s*(?:openai|anthropic|google)/[\w./-]+", "", t)
    t = re.sub(r"(?:\s*·\s*){2,}", " · ", t)
    t = re.sub(r"\((?:Killer II,\s*)?(?:openai|anthropic)/[\w./-]+\)", "", t)
    t = t.replace("Killer", "revisión crítica")
    # Etiquetas de afirmaciones y de supuestos.
    m = re.match(r"^(\s*)-\s*\[([^\]]*)\]\s*(.*)$", t)
    if m:
        dentro = m.group(2).lower()
        if dentro in _ESTADO_SUPUESTO:
            # La nota del evaluador va entre paréntesis al final: es su razonamiento.
            texto = re.sub(r"\s*\((?:[^()]|\([^()]*\))*\)\s*$", "", m.group(3)) if m.group(3).rstrip().endswith(")") else m.group(3)
            t = f"{m.group(1)}- {texto} ({_ESTADO_SUPUESTO[dentro]})"
        elif "," in dentro or dentro in ("sostenida", "parcial"):
            t = f"{m.group(1)}- " + ("(datos sintéticos) " if "sintetico" in dentro or "sintético" in dentro else "") + m.group(3)
    t = re.sub(r"^\s*-\s*([a-z]+(?:_[a-z]+)+)\s*$", lambda g: "- " + _BLOQUEO_LEGIBLE.get(g.group(1), g.group(1).replace("_", " ")), t)
    # Claves en snake_case sueltas: entre paréntesis sobran; en el texto, con espacios.
    t = re.sub(r"\s*\((?:[a-z]+_)+[a-z]+\)", "", t)
    t = re.sub(r"\b([a-z]+(?:_[a-z0-9]+)+)\b", lambda g: g.group(1).replace("_", " "), t)
    return re.sub(r"\s{2,}", " ", t).rstrip(" ,;")


def para_lector(contenido: str) -> str:
    """El dossier sin lo que solo sirve para auditarlo."""
    salida: list[str] = []
    fuera_hasta: int | None = None
    for x in (contenido or "").splitlines():
        cab = re.match(r"^(#+)\s+(.*)$", x)
        if cab:
            nivel = len(cab.group(1))
            if fuera_hasta is not None and nivel <= fuera_hasta:
                fuera_hasta = None
            if any(cab.group(2).lower().startswith(f) for f in _SECCIONES_FUERA):
                fuera_hasta = nivel
                continue
            if cab.group(2).lower().startswith("riesgo de sesgo por instrumento"):
                x = f"{cab.group(1)} Riesgo de sesgo de las fuentes"
        if fuera_hasta is not None or _LINEAS_FUERA.search(x):
            continue
        # El "pasaje" de un análisis es su salida en bruto (clave=valor): la cifra que
        # importa ya va en la afirmación.
        if re.match(r"^\s*Pasaje literal:", x) and x.count("=") >= 3:
            continue
        # Las comprobaciones una a una de una decisión: su motivo ya va en la línea.
        if re.match(r"^\s{2,}-\s", x):
            continue
        limpia = _limpiar_linea(x)
        if limpia.strip() in ("-", ""):
            continue
        salida.append(limpia)
    return "\n".join(salida)


# -- El PDF ---------------------------------------------------------------------------
#
# El formato que pidió Emir (25 de septiembre de 2026): el PDF, como el de la
# norma de Monica. Las medidas salen de su propio PDF, leídas con PyMuPDF: A4,
# márgenes de 72 puntos; logo en (73,5; 37,5) a (127,8; 72,2); cabecera en
# negrita 8,5 pt #3B145F alineada a la derecha; título en negrita 20 pt
# #3B145F; subtítulo 11 pt #613D8F; líneas de datos en negrita 9 pt #3B145F
# sobre una franja #F3ECFB; secciones en negrita 13 pt #3B145F; cuerpo 10,5 pt
# #1E2126; pie centrado 8 pt #613D8F. La fuente es la suya, Liberation Sans (la
# versión libre de Arial, licencia SIL OFL, en rosa/plantillas/fuentes), y el
# logo es la imagen que lleva su PDF.

FUENTES = Path(__file__).resolve().parent / "plantillas" / "fuentes"
LOGO_PDF = Path(__file__).resolve().parent / "plantillas" / "logo_alzheimer_project.jpg"
_A4 = (595.3, 841.9)
_MARGEN = 72.0
_CSS = """
@font-face { font-family: lib; src: url(LiberationSans-Regular.ttf); }
@font-face { font-family: lib; src: url(LiberationSans-Bold.ttf); font-weight: bold; }
body { font-family: lib; font-size: 10.5pt; color: #1E2126; }
p { margin: 0 0 4pt 0; line-height: 1.35; }
h1 { font-size: 20pt; font-weight: bold; color: #3B145F; margin: 0 0 6pt 0; line-height: 1.2; }
p.sub { font-size: 11pt; color: #613D8F; margin: 0 0 18pt 0; }
p.dato { font-size: 9pt; font-weight: bold; color: #3B145F; margin: 0; padding: 4pt 0 4pt 4pt; }
h2 { font-size: 13pt; font-weight: bold; color: #3B145F; margin: 16pt 0 6pt 0; }
p.vineta { margin: 0 0 3pt 0; }
"""


def _html_seguro(t: str) -> str:
    import html

    partes = re.split(r"\*\*", html.escape(t))
    return "".join(f"<b>{x}</b>" if i % 2 else x for i, x in enumerate(partes))


def html_del_dossier(contenido: str, id_doc: str | None, v: dict[str, Any]) -> str:
    """El cuerpo del documento en HTML, con las mismas partes que el Word:
    título, subtítulo, el bloque de datos de control y el dossier."""
    lineas = [x.rstrip() for x in para_lector(contenido).splitlines()]
    titulo = next((x[2:].strip() for x in lineas if x.startswith("# ")), "Dossier para el laboratorio")
    emitido = v.get("emitido", True)
    datos = [
        f"Document Status: Draft{'' if emitido else ' (no emitido como documento controlado)'}",
        f"Document ID: {id_doc if emitido and id_doc else 'sin asignar (se da al emitir)'}",
        f"Version: {v['version'][1:] if emitido else 'borrador'}",
        f"Date: {v['fecha']}",
        f"Owner Initials: {v['iniciales']}",
    ]
    h = [f"<h1>{_html_seguro(titulo)}</h1>", '<p class="sub">Dossier para el laboratorio, generado por ROSA2018</p>']
    h += [f'<p class="dato">{_html_seguro(x)}</p>' for x in datos]
    h.append('<p style="margin-bottom:8pt"> </p>')
    vio_titulo = False
    for x in lineas:
        if not x.strip():
            continue
        if x.startswith("# ") and not vio_titulo:
            vio_titulo = True
            continue
        if x.startswith("#"):
            h.append(f"<h2>{_html_seguro(x.lstrip('#').strip())}</h2>")
        elif re.match(r"^\s*[-*] ", x):
            h.append(f'<p class="vineta">- {_html_seguro(re.sub(r"^\s*[-*] ", "", x))}</p>')
        else:
            h.append(f"<p>{_html_seguro(x)}</p>")
    return "<body>" + "\n".join(h) + "</body>"


def pdf(contenido: str, id_doc: str | None, v: dict[str, Any]) -> bytes:
    """El dossier en PDF con el diseño de la norma AP-DOC-002. `v` es una
    versión emitida (con su código `id_doc`) o un `borrador`. El cuerpo se
    maqueta con la Story de PyMuPDF y después se pintan en cada página el logo,
    la cabecera y el pie, con "Page X of Y" ya contado."""
    import pymupdf

    ancho, alto = _A4
    pagina = pymupdf.Rect(0, 0, ancho, alto)
    donde = pymupdf.Rect(_MARGEN, 95, ancho - _MARGEN, alto - 75)
    story = pymupdf.Story(html=html_del_dossier(contenido, id_doc, v), user_css=_CSS, archive=pymupdf.Archive(str(FUENTES)))
    salida = io.BytesIO()
    escritor = pymupdf.DocumentWriter(salida)
    mas = True
    while mas:
        disp = escritor.begin_page(pagina)
        mas, _ = story.place(donde)
        story.draw(disp)
        escritor.end_page()
    escritor.close()

    doc = pymupdf.open("pdf", salida.getvalue())
    regular = pymupdf.Font(fontfile=str(FUENTES / "LiberationSans-Regular.ttf"))
    negrita = pymupdf.Font(fontfile=str(FUENTES / "LiberationSans-Bold.ttf"))
    morado, lila = (0x3B / 255, 0x14 / 255, 0x5F / 255), (0x61 / 255, 0x3D / 255, 0x8F / 255)
    total = len(doc)
    prefijo = v["pie"].split("Page X of Y")[0]
    # La franja lila del bloque de datos, solo en la primera página y por debajo
    # del texto. Con el fondo en el CSS, la Story lo volvía a pintar arriba de
    # todas las páginas siguientes.
    if total:
        p0 = doc[0]
        for linea in html_del_dossier(contenido, id_doc, v).split("\n"):
            m = re.match(r'<p class="dato">([^<:]+:)', linea)
            if not m:
                continue
            for r in p0.search_for(m.group(1))[:1]:
                p0.draw_rect(pymupdf.Rect(_MARGEN, r.y0 - 5, ancho - _MARGEN, r.y1 + 5.2), color=None, fill=(0xF3 / 255, 0xEC / 255, 0xFB / 255), overlay=False)
    for i in range(total):
        p = doc[i]
        if LOGO_PDF.exists():
            p.insert_image(pymupdf.Rect(73.5, 37.45, 127.85, 72.15), filename=str(LOGO_PDF))
        cab = str(v["cabecera"])
        w = negrita.text_length(cab, fontsize=8.5)
        tw = pymupdf.TextWriter(pagina)
        tw.append((ancho - _MARGEN - w, 64.5), cab, font=negrita, fontsize=8.5)
        tw.write_text(p, color=morado)
        texto_pie = f"{prefijo}Page {i + 1} of {total}"
        w = regular.text_length(texto_pie, fontsize=8)
        tw = pymupdf.TextWriter(pagina)
        tw.append(((ancho - w) / 2, 787.5), texto_pie, font=regular, fontsize=8)
        tw.write_text(p, color=lila)
    doc.set_metadata({"title": str(v.get("nombre") or ""), "author": "Alzheimer Project", "subject": f"{id_doc or 'Borrador'} {v.get('version')}", "creator": "ROSA2018"})
    return doc.tobytes(garbage=3, deflate=True)
