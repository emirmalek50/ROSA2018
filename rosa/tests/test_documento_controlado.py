"""Documentos de hipótesis controlados según la norma AP-DOC-002 v01
(rosa/documento_controlado.py): código AP-HYP correlativo que no se reutiliza,
versión v01, v02..., cabecera y pie con el formato exacto, las ocho
comprobaciones de la norma, el Word con la numeración como campo automático, la
acción y la descarga. Sin red ni modelos.
"""

import asyncio
import io
from types import SimpleNamespace

import pytest

from rosa import documento_controlado as DC
from rosa.estado import acciones as A

# 17 de septiembre de 2026 a mediodía, la fecha de la propia norma.
AHORA = 1789660800000


def _estado(n_hip: int = 2) -> dict:
    hips, arts = [], []
    for i in range(1, n_hip + 1):
        arts.append({"id": f"art-{i}", "investigacionId": "inv", "nombre": f"Dossier {i}", "tipo": "dossier", "destacado": False, "versiones": [{"n": 1, "creadaEn": 1, "resumen": "", "contenido": f"# Dossier {i}\n\n## 1. Decisión\n- Un **punto**\nTexto.", "iteracion": 1, "procedencia": {}}]})
        hips.append({"id": f"h{i}", "investigacionId": "inv", "titulo": f"Título largo de la hipótesis {i}", "estado": "propuesta", "dossierArtefactoId": f"art-{i}", "procedencia": {"registro": []}})
    return {"hipotesis": hips, "artefactos": arts, "eventos": [], "investigaciones": [{"id": "inv"}]}


# -- Lo que fija la norma ------------------------------------------------------------


def test_formatos_de_la_norma():
    assert DC.id_documento(1) == "AP-HYP-001" and DC.id_documento(42) == "AP-HYP-042" and DC.id_documento(1234) == "AP-HYP-1234"
    assert DC.version_documento(1) == "v01" and DC.version_documento(12) == "v12"
    assert DC.fecha_documento(AHORA) == "Sep-17-2026"
    # Los mismos ejemplos que trae la norma.
    assert DC.cabecera("Amyloid-Tau Neuroinflammation Mechanism", "AP-HYP-001", "v01") == "Amyloid-Tau Neuroinflammation Mechanism | DOC | AP-HYP-001 | v01"
    assert DC.pie("Sep-17-2026", "MD") == "Confidential | Alzheimer Project | AI Robotix | Sep-17-2026 | MD | Page X of Y"


def test_el_nombre_corto_no_puede_romper_la_cabecera():
    assert DC.limpiar_nombre("  Brecha | GFAP  NfL.  ") == "Brecha GFAP NfL"
    assert DC.limpiar_nombre("«Brecha GFAP»") == "Brecha GFAP"
    assert DC.limpiar_nombre("a\nb") == "a b"
    assert DC.limpiar_nombre(None) == "" and DC.limpiar_nombre("   ") == ""
    assert len(DC.limpiar_nombre("x" * 500)) == DC.MAX_NOMBRE


# -- Emitir ------------------------------------------------------------------------------


def test_la_primera_emision_da_codigo_y_v01_y_pasa_las_ocho_comprobaciones():
    e = _estado()
    v = DC.emitir(e, e["hipotesis"][0], "Brecha GFAP NfL", "Emir", AHORA)
    assert isinstance(v, dict)
    assert e["hipotesis"][0]["documentoControlado"]["id"] == "AP-HYP-001"
    assert v["version"] == "v01" and v["fecha"] == "Sep-17-2026" and v["iniciales"] == DC.INICIALES_RESPONSABLE
    assert v["cabecera"] == "Brecha GFAP NfL | DOC | AP-HYP-001 | v01"
    assert v["pie"] == f"Confidential | Alzheimer Project | AI Robotix | Sep-17-2026 | {DC.INICIALES_RESPONSABLE} | Page X of Y"
    assert [c["n"] for c in v["comprobaciones"]] == list(range(1, 9))
    assert all(c["ok"] for c in v["comprobaciones"]) and v["controlado"] is True
    assert v["artefactoId"] == "art-1" and v["versionArtefacto"] == 1


def test_los_codigos_son_correlativos_y_no_se_reutilizan():
    e = _estado(3)
    DC.emitir(e, e["hipotesis"][0], "Uno", "E", AHORA)
    DC.emitir(e, e["hipotesis"][1], "Dos", "E", AHORA)
    assert [h.get("documentoControlado", {}).get("id") for h in e["hipotesis"]] == ["AP-HYP-001", "AP-HYP-002", None]
    # Aunque el contador se pierda (un estado restaurado a mano), no se repite.
    e.pop("documentosControlados")
    DC.emitir(e, e["hipotesis"][2], "Tres", "E", AHORA)
    assert e["hipotesis"][2]["documentoControlado"]["id"] == "AP-HYP-003"
    # Y aunque la hipótesis del 002 desaparezca, el contador lo recuerda.
    e["hipotesis"].pop(1)
    e["hipotesis"].append({"id": "h9", "investigacionId": "inv", "titulo": "t", "dossierArtefactoId": "art-2", "procedencia": {"registro": []}})
    DC.emitir(e, e["hipotesis"][-1], "Nueve", "E", AHORA)
    assert e["hipotesis"][-1]["documentoControlado"]["id"] == "AP-HYP-004"


def test_la_misma_version_del_dossier_no_se_emite_dos_veces_y_la_nueva_sube_version():
    e = _estado()
    h = e["hipotesis"][0]
    DC.emitir(e, h, "Uno", "E", AHORA)
    motivo = DC.emitir(e, h, "Uno", "E", AHORA)
    assert isinstance(motivo, str) and "ya está emitida como v01" in motivo
    # Se regenera el dossier: versión nueva del artefacto, v02 del documento, mismo código.
    e["artefactos"][0]["versiones"].append({"n": 2, "contenido": "# Dossier nuevo", "procedencia": {}})
    v2 = DC.emitir(e, h, None, "E", AHORA + 86_400_000)
    assert v2["version"] == "v02" and h["documentoControlado"]["id"] == "AP-HYP-001"
    assert v2["nombre"] == "Uno"  # sin nombre nuevo, se conserva el que tenía
    assert v2["fecha"] == "Sep-18-2026" and v2["versionArtefacto"] == 2
    # La v01 sigue diciendo lo que decía.
    assert DC.version_emitida(h, "v01")["versionArtefacto"] == 1 and DC.version_emitida(h, "v01")["fecha"] == "Sep-17-2026"


def test_sin_dossier_o_sin_nombre_no_se_emite():
    e = _estado()
    h = e["hipotesis"][0]
    h["dossierArtefactoId"] = None
    assert "genéralo primero" in DC.emitir(e, h, "Uno", "E", AHORA)
    h["dossierArtefactoId"] = "art-1"
    assert "Falta el nombre corto" in DC.emitir(e, h, "  |  ", "E", AHORA)
    assert "documentoControlado" not in h and "documentosControlados" not in e
    # El nombre que resumió ROSA2018 vale si la persona no escribe otro.
    h["nombreCorto"] = "Resumido por ROSA2018"
    assert DC.emitir(e, h, "", "E", AHORA)["nombre"] == "Resumido por ROSA2018"


def test_las_comprobaciones_cazan_un_documento_que_no_cumple():
    v = {"version": "v1", "fecha": "17/09/2026", "iniciales": "", "nombre": "x", "cabecera": "x | DOC | AP-HYP-001 | v1", "pie": "Confidential | Alzheimer Project | AI Robotix | 17/09/2026 |  | Page X of Y"}
    fallan = {c["n"] for c in DC.comprobar("AP-HYP-1", v) if not c["ok"]}
    assert fallan == {1, 2, 5, 6, 7, 8}
    assert {c["n"] for c in DC.comprobar("XX-DOC-001", v) if not c["ok"]} >= {4, 5}


# -- La acción y el evento -----------------------------------------------------------------


def test_la_accion_deja_rastro_y_evento():
    e = _estado()
    assert A.emitir_documento(e, "h1", "Brecha GFAP NfL", "Emir", AHORA) is True
    h = e["hipotesis"][0]
    assert "documento controlado AP-HYP-001 v01 emitido por Emir" in h["procedencia"]["registro"][-1]
    assert any("Documento controlado AP-HYP-001 v01 emitido: Brecha GFAP NfL" in ev["texto"] for ev in e["eventos"])
    assert A.emitir_documento(e, "h1", "Brecha GFAP NfL", "Emir", AHORA) is False  # la misma versión
    assert A.emitir_documento(e, "no-existe", "x", "Emir", AHORA) is False


def test_la_accion_esta_registrada_para_el_servidor():
    from rosa.estado.almacen import _TABLA

    assert _TABLA["emitirDocumento"] is A.emitir_documento


# -- El Word ---------------------------------------------------------------------------------


def test_el_word_lleva_el_diseno_de_la_norma_la_cabecera_el_pie_con_campos_y_el_cuerpo():
    import docx

    e = _estado()
    v = DC.emitir(e, e["hipotesis"][0], "Brecha GFAP NfL", "E", AHORA)
    d = docx.Document(io.BytesIO(DC.docx(e["artefactos"][0]["versiones"][0]["contenido"], "AP-HYP-001", v)))
    s = d.sections[0]
    # La cabecera de la plantilla de la norma: el logo y la línea de control.
    assert v["cabecera"] in [p.text for p in s.header.paragraphs]
    assert "graphicData" in s.header._element.xml  # el logo del Alzheimer Project
    pie = s.footer.paragraphs[0]
    assert pie.text.startswith(f"Confidential | Alzheimer Project | AI Robotix | Sep-17-2026 | {DC.INICIALES_RESPONSABLE} | Page ")
    xml = pie._p.xml
    # La numeración es un campo de Word, nunca un número escrito a mano.
    assert "PAGE" in xml and "NUMPAGES" in xml
    textos = [p.text for p in d.paragraphs]
    assert textos[0] == "Dossier 1" and "Document ID: AP-HYP-001" in textos and "Version: 01" in textos
    assert "1. Decisión" in textos and "- Un punto" in textos
    assert not any("{{" in t for t in textos)  # no queda ninguna muestra de la plantilla
    # Los colores de la norma: el título en #3B145F, en negrita.
    titulo = d.paragraphs[0].runs[0]
    assert str(titulo.font.color.rgb) == "3B145F" and titulo.bold
    negritas = [r.text for p in d.paragraphs for r in p.runs if r.bold]
    assert "punto" in negritas
    # Arial explícito en todo el cuerpo.
    assert all(r.font.name == "Arial" for p in d.paragraphs for r in p.runs)


def test_un_dossier_sin_emitir_se_descarga_como_borrador_sin_codigo():
    import docx

    e = _estado()
    contenido, id_doc, v = DC.para_descargar(e, "art-1", 1, AHORA)
    assert id_doc is None and v["emitido"] is False and "sin código: no emitido | borrador" in v["cabecera"]
    d = docx.Document(io.BytesIO(DC.docx(contenido, id_doc, v)))
    textos = [p.text for p in d.paragraphs]
    assert "Document ID: sin asignar (se da al emitir)" in textos and any("no emitido" in t for t in textos)
    # Emitido, la misma versión sale con su código; otra versión, borrador.
    DC.emitir(e, e["hipotesis"][0], "Brecha", "E", AHORA)
    assert DC.para_descargar(e, "art-1", 1, AHORA)[1] == "AP-HYP-001"
    e["artefactos"][0]["versiones"].append({"n": 2, "contenido": "# Nuevo", "procedencia": {}})
    assert DC.para_descargar(e, "art-1", 2, AHORA)[1] is None
    assert DC.para_descargar(e, "art-1", 3, AHORA) is None and DC.para_descargar(e, "no", 1, AHORA) is None


# -- La descarga -----------------------------------------------------------------------------


def test_la_descarga_sirve_la_version_emitida_y_no_otra(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from rosa import servidor as S
    from rosa.estado.almacen import Almacen

    al = Almacen(str(tmp_path / "rosa.db"))

    def preparar(e):
        base = _estado()
        e["hipotesis"] = base["hipotesis"]
        e["artefactos"] = base["artefactos"]
        return True

    al.mutar(preparar, "t")
    al.mutar(lambda e: A.emitir_documento(e, "h1", "Brecha GFAP NfL", "E", AHORA), "t")
    app = S.crear_app(al)
    app.state.acceso = SimpleNamespace(usuario=lambda token: "emir@alzheimerproject.com" if token == "sesion-test" else None, es_admin=lambda email: False, salir=lambda token: None)
    c = TestClient(app, base_url="http://127.0.0.1:8765", cookies={"rosa_sesion": "sesion-test"})
    r = c.get("/api/documentos/h1/v01.pdf")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/pdf")
    assert 'filename="AP-HYP-001_v01.pdf"' in r.headers["content-disposition"]
    assert c.get("/api/documentos/h1/v02.pdf").status_code == 404
    assert c.get("/api/documentos/nadie/v01.pdf").status_code == 404
    r = c.get("/api/artefactos/art-1/v/1.pdf")
    assert r.status_code == 200 and 'filename="AP-HYP-001_v01.pdf"' in r.headers["content-disposition"]
    assert c.get("/api/artefactos/art-2/v/1.pdf").headers["content-disposition"].endswith('Dossier_borrador_v1.pdf"')
    assert c.get("/api/artefactos/art-1/v/9.pdf").status_code == 404
    # Sin sesión no se descarga: el documento es confidencial.
    anonimo = TestClient(app, base_url="http://127.0.0.1:8765")
    assert anonimo.get("/api/documentos/h1/v01.pdf").status_code == 401
    assert anonimo.get("/api/artefactos/art-1/v/1.pdf").status_code == 401


# -- El nombre corto que resume ROSA2018 -------------------------------------------------------


def _supervisor(estado, llamar):
    from rosa.bucle import corrida as COR

    sup = COR.Supervisor.__new__(COR.Supervisor)
    guardado = {"e": estado}

    def mutar(fn, _motivo):
        fn(guardado["e"])

    sup.almacen = SimpleNamespace(estado=estado, mutar=mutar)
    sup.programas = SimpleNamespace(nombre_corto="nombre_corto")
    return sup, SimpleNamespace(llamar=llamar)


def test_el_nombre_corto_se_resume_y_se_limpia():
    e = _estado(1)

    async def llamar(rol, programa, **kw):
        assert rol == "volumen" and programa == "nombre_corto" and kw["titulo"] == "Título largo de la hipótesis 1"
        return SimpleNamespace(nombre='"Brecha | GFAP NfL."')

    sup, ctx = _supervisor(e, llamar)
    asyncio.run(sup._nombre_corto(ctx, e["hipotesis"][0]))
    assert e["hipotesis"][0]["nombreCorto"] == "Brecha GFAP NfL" and e["hipotesis"][0]["_nombreCortoIntentado"] is True


def test_un_nombre_demasiado_largo_o_sin_presupuesto_queda_para_la_persona():
    from rosa.modulos.contador import PresupuestoAgotado

    e = _estado(1)

    async def largo(rol, programa, **kw):
        return SimpleNamespace(nombre=" ".join(["palabra"] * 15))

    sup, ctx = _supervisor(e, largo)
    asyncio.run(sup._nombre_corto(ctx, e["hipotesis"][0]))
    assert e["hipotesis"][0]["nombreCorto"] is None and e["hipotesis"][0]["_nombreCortoIntentado"] is True

    e = _estado(1)

    async def sin(rol, programa, **kw):
        raise PresupuestoAgotado("sin presupuesto")

    sup, ctx = _supervisor(e, sin)
    asyncio.run(sup._nombre_corto(ctx, e["hipotesis"][0]))
    assert e["hipotesis"][0]["nombreCorto"] is None and e["hipotesis"][0]["_nombreCortoIntentado"] is True


# -- El PDF, el formato que se descarga ------------------------------------------------


def test_el_pdf_lleva_el_diseno_de_la_norma_en_cada_pagina():
    import pymupdf

    e = _estado()
    largo = "# Dossier 1\n\n" + "\n".join(f"## {i}. Sección\n" + "Texto con Δ y ε y tildes: señal, está. " * 30 for i in range(1, 9))
    e["artefactos"][0]["versiones"][0]["contenido"] = largo
    v = DC.emitir(e, e["hipotesis"][0], "Brecha GFAP NfL", "E", AHORA)
    d = pymupdf.open("pdf", DC.pdf(largo, "AP-HYP-001", v))
    assert len(d) >= 2
    for i, p in enumerate(d):
        texto = p.get_text()
        assert "Brecha GFAP NfL | DOC | AP-HYP-001 | v01" in texto
        assert f"Confidential | Alzheimer Project | AI Robotix | Sep-17-2026 | {DC.INICIALES_RESPONSABLE} | Page {i + 1} of {len(d)}" in texto
        assert p.get_images(), "falta el logo del Alzheimer Project"
    primera = d[0].get_text("dict")["blocks"]
    spans = [s for b in primera for l in b.get("lines", []) for s in l["spans"]]
    titulo = next(s for s in spans if s["text"].startswith("Dossier 1"))
    assert round(titulo["size"]) == 20 and titulo["color"] == 0x3B145F and "Bold" in titulo["font"]
    assert all("LiberationSans" in s["font"] for s in spans)
    todo = "".join(p.get_text() for p in d)
    assert "Δ" in todo and "ε" in todo and "señal" in todo and "Document ID: AP-HYP-001" in todo
    assert d.metadata["author"] == "Alzheimer Project"


def test_el_pdf_de_un_borrador_lo_dice_y_no_da_codigo():
    import pymupdf

    e = _estado()
    contenido, id_doc, v = DC.para_descargar(e, "art-1", 1, AHORA)
    texto = pymupdf.open("pdf", DC.pdf(contenido, id_doc, v))[0].get_text()
    assert "sin código: no emitido | borrador" in texto and "Document ID: sin asignar" in texto and "AP-HYP-" not in texto
